"""產生「評測題庫」空白範本 Excel，給你填 50 題的問題與標準答案。

用法：
    python -m scripts.make_eval_template            # 產生 data/eval/eval_questions.xlsx
    python -m scripts.make_eval_template --rows 50  # 指定預留幾列（預設 50）

填法：
    - 「問題」：病人／家屬會問的衛教問題（必填）。
    - 「參考答案（標準答案）」：你認為理想的正確回答，之後拿來和模型回答並排比對。
    - 「題號」已預先填好 1..N，可不用動。
填完存檔後，執行：  python -m scripts.run_eval
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from app import config  # noqa: E402

# 欄位（順序即 Excel 欄位順序）；run_eval.py 會依這些標題讀取。
# 「衛教主題」可留空；有填的話結果會多一份「分主題轉介統計」。
COLUMNS = ["題號", "衛教主題", "問題", "參考答案（標準答案）"]
COL_WIDTHS = [8, 12, 52, 66]

HEADER_FILL = PatternFill("solid", fgColor="2F5597")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)


def main() -> None:
    parser = argparse.ArgumentParser(description="產生評測題庫空白範本 Excel")
    parser.add_argument("--rows", type=int, default=50, help="預留幾列（預設 50）")
    parser.add_argument(
        "--output",
        default=str(config.DATA_DIR / "eval" / "eval_questions.xlsx"),
        help="輸出檔路徑",
    )
    args = parser.parse_args()

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        ans = input(f"{out} 已存在，覆蓋會清掉你填過的內容。確定覆蓋？(y/N) ")
        if ans.strip().lower() != "y":
            print("已取消。")
            return

    wb = Workbook()
    ws = wb.active
    ws.title = "題庫"

    # 標題列
    for c, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = COL_WIDTHS[c - 1]

    # 在標題加填寫說明（欄序：題號 / 衛教主題 / 問題 / 參考答案）
    ws.cell(row=1, column=2).comment = Comment(
        "可留空。若填（例如「防噎」「防跌倒」），結果會多一份「分主題轉介統計」。",
        "範本",
    )
    ws.cell(row=1, column=3).comment = Comment(
        "填病人／家屬會問的衛教問題，一列一題。\n"
        "空白列會被略過，可只填你要測的題數。",
        "範本",
    )
    ws.cell(row=1, column=4).comment = Comment(
        "填你認為理想的正確答案；跑完評測後會和模型回答並排，供個管師比對打分。",
        "範本",
    )

    # 預先填好題號，其餘留白
    for i in range(1, args.rows + 1):
        r = i + 1
        ws.cell(row=r, column=1, value=i).alignment = Alignment(horizontal="center")
        for c in (2, 3, 4):
            ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"  # 捲動時固定標題列

    wb.save(out)
    print(f"已建立範本：{out}")
    print(f"  - 已預留 {args.rows} 列，題號 1..{args.rows} 已填好。")
    print("  - 請填「問題」與「參考答案（標準答案）」，存檔後執行： python -m scripts.run_eval")


if __name__ == "__main__":
    main()
