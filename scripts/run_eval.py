"""讀「評測題庫」→ 逐題跑 RAG 問答 → 輸出可打分的結果 Excel。

流程與 /ask 端點完全相同（store.query → llm.answer_question），
所以評測結果就是產品的真實行為。評測期間**不會**寫入正式轉介清單。

用法：
    python -m scripts.run_eval                       # 讀 data/eval/eval_questions.xlsx
    python -m scripts.run_eval --input 我的題庫.xlsx  # 指定題庫（支援 .xlsx / .csv）
    python -m scripts.run_eval --output 結果.xlsx     # 指定輸出檔

輸出：data/eval/eval_results_<模型>_<時間>.xlsx
  - 前半：題號 / 問題 / 參考答案 / 模型回答 / 是否轉介 / 引用出處 / 檢索到的來源
  - 後半（黃色欄）：留給個案管理師打分——準確度、有依據、有用性、該不該轉介、備註
  - 另一張「統計摘要」：模型、轉介率等可自動算的指標
"""
from __future__ import annotations

import argparse
import csv
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from app import config, llm, store  # noqa: E402

# 結果欄位；標「(填)」者為留給個管師打分的黃色欄。
RESULT_COLUMNS = [
    ("題號", 8),
    ("衛教主題", 12),
    ("問題", 42),
    ("參考答案（標準答案）", 52),
    ("模型回答", 52),
    ("是否轉介（模型）", 14),
    ("引用出處", 42),
    ("檢索到的來源", 26),
    ("準確度(1-5)(填)", 12),
    ("有依據(是/否)(填)", 13),
    ("有用性(1-5)(填)", 12),
    ("該不該轉介(是/否)(填)", 16),
    ("備註(填)", 30),
]
N_SCORE_START = 9  # 第 9 欄（準確度）起為留給個管師的打分欄
COL_ESC = 6        # 「是否轉介（模型）」所在欄

HEADER_FILL = PatternFill("solid", fgColor="2F5597")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
SCORE_FILL = PatternFill("solid", fgColor="FFF2CC")  # 打分欄底色（淺黃）
ESC_FONT = Font(color="C00000", bold=True)  # 轉介標紅


def _read_questions(path: Path) -> list[dict]:
    """從 .xlsx / .csv 讀題庫，回傳 [{id, question, reference}]。空白題目略過。"""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            rows = [list(r) for r in csv.reader(f)]
    else:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = [[c if c is not None else "" for c in r]
                for r in ws.iter_rows(values_only=True)]

    if not rows:
        return []

    header = [str(h or "").strip() for h in rows[0]]

    def find(*keys, default=None):
        for i, h in enumerate(header):
            if any(k in h for k in keys):
                return i
        return default

    i_id = find("題號", "序號")
    i_topic = find("主題")
    i_q = find("問題", "題目")
    i_ref = find("參考", "標準答案", "答案")
    if i_q is None:
        raise SystemExit(f"在 {path} 的第一列找不到「問題」欄，請用 make_eval_template 產生的範本。")

    items: list[dict] = []
    for n, row in enumerate(rows[1:], start=1):
        def cell(idx):
            return str(row[idx]).strip() if idx is not None and idx < len(row) else ""

        q = cell(i_q)
        if not q:
            continue  # 空白列略過
        items.append({
            "id": cell(i_id) or str(n),
            "topic": cell(i_topic),
            "question": q,
            "reference": cell(i_ref),
        })
    return items


def _run_one(question: str) -> dict:
    """跑單題，回傳結果欄位。任何例外都吞成一列錯誤，不中斷整批。"""
    try:
        hits = store.query(question)
        ans = llm.answer_question(question, hits)
        cites = []
        for idx in ans.sources_used:
            if 1 <= idx <= len(hits):
                h = hits[idx - 1]
                cites.append(f"[{idx}] {h['title']}：{h['text'][:120]}")
        retrieved = "、".join(dict.fromkeys(h["title"] for h in hits)) if hits else "（無檢索結果）"
        return {
            "answer": ans.answer,
            "escalated": not ans.can_answer,
            "citations": "\n".join(cites),
            "retrieved": retrieved,
            "error": "",
        }
    except Exception as e:  # noqa: BLE001
        return {
            "answer": "",
            "escalated": True,
            "citations": "",
            "retrieved": "",
            "error": f"{type(e).__name__}: {e}",
        }


def _write_results(out: Path, items: list[dict], results: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "評測結果"

    # 標題列
    for c, (name, width) in enumerate(RESULT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "C2"

    for r, (item, res) in enumerate(zip(items, results), start=2):
        vals = [
            item["id"],
            item.get("topic", ""),
            item["question"],
            item["reference"],
            res["error"] and f"⚠ 錯誤：{res['error']}" or res["answer"],
            "是" if res["escalated"] else "否",
            res["citations"],
            res["retrieved"],
            "", "", "", "", "",  # 打分欄留白
        ]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if c >= N_SCORE_START:
                cell.fill = SCORE_FILL
            if c == COL_ESC and res["escalated"]:
                cell.font = ESC_FONT
                cell.alignment = Alignment(horizontal="center", vertical="top")

    # 打分欄下拉選單，方便個管師勾選（欄位對應 RESULT_COLUMNS 順序）
    last = len(items) + 1
    dv_15 = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    dv_yn = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    ws.add_data_validation(dv_15)
    ws.add_data_validation(dv_yn)
    dv_15.add(f"I2:I{last}")   # 準確度(1-5)
    dv_15.add(f"K2:K{last}")   # 有用性(1-5)
    dv_yn.add(f"J2:J{last}")   # 有依據(是/否)
    dv_yn.add(f"L2:L{last}")   # 該不該轉介(是/否)

    _write_summary(wb, items, results)
    wb.save(out)


def _write_summary(wb: Workbook, items: list[dict], results: list[dict]) -> None:
    ws = wb.create_sheet("統計摘要")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60

    n = len(items)
    n_esc = sum(1 for r in results if r["escalated"])
    n_err = sum(1 for r in results if r["error"])
    answered = [r for r in results if not r["escalated"] and not r["error"]]
    avg_len = round(sum(len(r["answer"]) for r in answered) / len(answered)) if answered else 0
    docs = "、".join(f"{d['title']}({d['chunks']}片段)" for d in store.list_documents()) or "（空）"

    rows = [
        ("評測時間", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("LLM 供應商", config.LLM_PROVIDER),
        ("LLM 模型", config.LLM_MODEL),
        ("溫度", config.LLM_TEMPERATURE),
        ("檢索 top_k", config.TOP_K),
        ("知識庫文件", docs),
        ("", ""),
        ("總題數", n),
        ("模型判定可回答", len(answered)),
        ("模型判定轉介", n_esc),
        ("轉介率", f"{n_esc / n:.0%}" if n else "-"),
        ("執行錯誤題數", n_err),
        ("可答題平均字數", avg_len),
        ("", ""),
        ("※ 準確率／有用性", "請個管師在「評測結果」分頁填完評分後再統計"),
        ("※ 轉介判斷正確率", "= 模型「是否轉介」與個管師「該不該轉介」一致的比例"),
    ]
    for r, (k, v) in enumerate(rows, start=1):
        a = ws.cell(row=r, column=1, value=k)
        a.font = Font(bold=True)
        ws.cell(row=r, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")

    # 分主題統計（若題庫有填「衛教主題」）
    topics = [it.get("topic", "") for it in items]
    if any(topics):
        r = len(rows) + 2
        ws.cell(row=r, column=1, value="── 分主題轉介統計 ──").font = Font(bold=True)
        r += 1
        ws.cell(row=r, column=1, value="衛教主題").font = Font(bold=True)
        ws.cell(row=r, column=2, value="題數 / 轉介數 / 轉介率").font = Font(bold=True)
        seen: list[str] = []
        for t in topics:
            key = t or "（未分類）"
            if key in seen:
                continue
            seen.append(key)
            idxs = [i for i, tt in enumerate(topics) if (tt or "（未分類）") == key]
            total = len(idxs)
            esc = sum(1 for i in idxs if results[i]["escalated"])
            r += 1
            ws.cell(row=r, column=1, value=key).font = Font(bold=True)
            ws.cell(row=r, column=2, value=f"{total} / {esc} / {esc / total:.0%}")


def main() -> None:
    parser = argparse.ArgumentParser(description="跑評測題庫，輸出可打分的結果 Excel")
    parser.add_argument(
        "--input",
        default=str(config.DATA_DIR / "eval" / "eval_questions.xlsx"),
        help="題庫檔（.xlsx 或 .csv），預設 data/eval/eval_questions.xlsx",
    )
    parser.add_argument("--output", default=None, help="輸出檔路徑（預設自動命名）")
    args = parser.parse_args()

    in_path = Path(args.input)
    if not in_path.exists():
        raise SystemExit(f"找不到題庫檔：{in_path}\n先執行： python -m scripts.make_eval_template")

    items = _read_questions(in_path)
    if not items:
        raise SystemExit(f"{in_path} 沒有讀到任何題目（請確認「問題」欄已填）。")

    if args.output:
        out_path = Path(args.output)
    else:
        safe_model = config.LLM_MODEL.split("/")[-1].replace(":", "_")
        stamp = datetime.now().strftime("%Y%m%d_%H%M")
        out_path = config.DATA_DIR / "eval" / f"eval_results_{safe_model}_{stamp}.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)

    print(f"題庫：{in_path}（{len(items)} 題）")
    print(f"模型：{config.LLM_PROVIDER} / {config.LLM_MODEL}")
    print(f"知識庫：{'、'.join(d['title'] for d in store.list_documents()) or '（空）'}\n")

    results = []
    for i, item in enumerate(items, start=1):
        print(f"[{i}/{len(items)}] {item['question'][:30]}...", flush=True)
        res = _run_one(item["question"])
        flag = "⚠錯誤" if res["error"] else ("→轉介" if res["escalated"] else "已回答")
        print(f"        {flag}", flush=True)
        results.append(res)

    _write_results(out_path, items, results)
    n_esc = sum(1 for r in results if r["escalated"])
    print(f"\n完成！{len(items)} 題，其中 {n_esc} 題判定轉介。")
    print(f"結果已存：{out_path}")
    print("請把此檔交給個案管理師，填黃色欄（準確度／有依據／有用性／該不該轉介）。")


if __name__ == "__main__":
    main()
